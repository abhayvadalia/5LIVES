from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
out=Path(__file__).resolve().parents[1]/'Five_Lives_Simple_Explainer.docx'
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(.75)
s.left_margin=s.right_margin=Inches(.85)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
st=doc.styles['Normal']; st.font.size=Pt(12); st.paragraph_format.space_after=Pt(9); st.paragraph_format.line_spacing=1.1
st=doc.styles['Title']; st.font.size=Pt(28); st.paragraph_format.space_after=Pt(14)
st=doc.styles['Heading 1']; st.font.size=Pt(16); st.paragraph_format.space_before=Pt(15); st.paragraph_format.space_after=Pt(7)
def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t,1)
doc.add_paragraph('Five Lives explained simply','Title')
p('Five Lives is an idea for a platform that helps people imagine five different ways their life could look, then try small parts of those lives in the real world. The aim is to help people learn what suits them before making bigger decisions.')
h('Why someone would use it')
p('A person can have more than one dream. They might enjoy their current life and still wonder what it would be like to run a café, make music, travel more, teach, or spend more time outdoors. It can be difficult to know which interests deserve more time when they exist only in the imagination.')
p('Five Lives gives those possibilities a practical starting point. Its central belief is that a person can experience several ways of living within one lifetime. This might mean a new hobby, a different routine, a side project, or eventually a larger change.')
h('What a possible life means')
p('A possible life is a picture of how you would spend your time and what would matter to you. It can include work, relationships, surroundings, interests, and everyday routines. It does not have to mean changing your career or starting again.')
p('For example, a creative life might mean drawing for an hour each evening. An adventurous life might mean exploring a new trail on weekends. The important question is what you would actually do and whether you would enjoy doing it regularly.')
h('How the experience could work')
p('The first version will focus on each person’s own journey. A proposed approach is to describe five possible lives, look at the everyday reality of each one, and choose a small activity to try for a limited time. Suggested trials of 7 to 30 days and check-ins are still being considered.')
p('After trying an activity, you could reflect on what felt enjoyable, what was difficult, and what you would like to keep doing. The intended result is a clearer next step based on experience. The exact screens and activities are still to be designed.')
doc.add_page_break()
h('A simple example')
p('Imagine Maya, who has a regular office job and several interests. She writes down five lives she is curious about: café owner, photographer, outdoor explorer, teacher, and musician. These are examples, not fixed categories everyone would have to choose.')
p('Maya starts with the café idea. Instead of judging it only by the appealing image of a cosy shop, she explores the daily work. With a café owner’s agreement, she could spend a few mornings observing preparation, service, and cleaning. She could also think through the time and costs involved.')
p('She might discover that she loves making food and talking to customers but dislikes the early starts and repetitive tasks. Her next step could be a weekend baking hobby. Or she might want to investigate the business further. Either result helps her understand what she actually wants.')
p('She can then explore another interest. She does not need to turn all five lives into full-time commitments. Parts of different lives could fit into the life she already has.')
h('Why it would feel like an exploration game')
p('The chosen direction is bright, playful, and polished, with plenty of motion. Exploring possibilities should feel inviting and adventurous. A colourful map with five regions has been suggested as a way to show the journey, but the final design is not settled.')
p('The activities would happen in real life. The game-like presentation would help people see where they are in their journey and feel encouraged to take the next step.')
h('Where the community fits in')
p('The longer-term idea is to connect people who discover similar interests or visions for their lives. Someone exploring photography, for example, could meet others interested in learning and practising it.')
p('This community is envisioned as a gated space, meaning access would be limited to members, with subscriptions supporting membership. It comes after the personal journey in the current priorities. Prices, membership options, and the details of how people would be connected are still open.')
h('What Five Lives aims to give people')
p('Five Lives aims to turn “I wonder what that life would be like” into something a person can explore. By trying manageable activities and reflecting on them, people could make more informed choices about what to bring into their everyday lives.')
p('The product is currently a concept in development. Its core direction is agreed, while the detailed experience and business model are still being worked out.')
doc.core_properties.title='Five Lives explained simply'
doc.core_properties.subject='A plain English explanation of the Five Lives concept'
doc.core_properties.author='Five Lives'
for el in list(doc.styles.element.iter()):
    for child in list(el):
        if child.tag.endswith('}pBdr'):
            el.remove(child)
for el in list(doc.element.iter()):
    for child in list(el):
        if child.tag.endswith('}pBdr'):
            el.remove(child)
doc.save(out)
print(out)
